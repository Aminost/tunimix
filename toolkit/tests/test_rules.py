"""Rule-based engine (no AI): parsing, step conversion, structuring, documents, OCR, pipeline."""
import io
import shutil

import pytest

from tunimix.convert.ingredients import detect_language, parse_ingredient
from tunimix.convert.rulebased import convert_structured
from tunimix.extract.documents import read_document, structure_text
from tunimix.formatting import step_text_with_setting

OJJA = {"title": "Ojja aux merguez", "yields": "4 personnes",
        "ingredients": ["300 g de merguez", "4 oeufs", "3 gousses d'ail", "2 c. à soupe de concentré de tomate",
                        "1 c. à café de harissa", "4 c. à soupe d'huile d'olive", "2 poivrons verts", "400 g de tomates", "sel, poivre"],
        "instructions": ["Hacher l'ail finement.", "Faire revenir l'ail dans l'huile d'olive, ajouter le concentré de tomate et la harissa.",
                         "Ajouter les poivrons et les tomates, puis laisser mijoter 15 minutes.",
                         "Ajouter les merguez coupées en rondelles et cuire 10 min.", "Casser les oeufs dans la sauce et laisser cuire 4 minutes.",
                         "Servir chaud avec du pain."]}
SOUP_TEXT = """Linsensuppe

Für 4 Personen

Zutaten
1 Zwiebel
2 Knoblauchzehen
2 EL Olivenöl
250 g rote Linsen
1 l Gemüsebrühe
1 TL Kreuzkümmel
½ Bund Petersilie

Zubereitung
1. Zwiebel und Knoblauch fein hacken und im Olivenöl
andünsten.
2. Linsen, Kreuzkümmel und Brühe dazugeben und 20 Minuten köcheln lassen.
3. Die Suppe fein pürieren.
4. Petersilie hacken und darüber streuen.
"""


def texts(r):
    return [step_text_with_setting(s.text, s.setting, r.language)[0] for s in r.steps]


@pytest.mark.parametrize("line,lang,exp", [
    ("2 EL Olivenöl, extra vergine", "de", ("2", "EL", "Olivenöl", "extra vergine")),
    ("1 cup flour (sifted)", "en", ("120", "g", "flour", "sifted")),
    ("3 gousses d'ail", "fr", ("3", "gousse(s)", "ail", "")),
    ("1 verre d'huile d'olive", "fr", ("200", "ml", "huile d'olive", "")),
    ("1-2 kg Kartoffeln", "de", ("1000-2000", "g", "Kartoffeln", "")),
    ("Salz, Pfeffer", "de", ("", "", "Salz, Pfeffer", "")),
])
def test_ingredient_parser(line, lang, exp):
    i = parse_ingredient(line, lang)
    assert (i.amount, i.unit, i.name, i.note) == exp


def test_language_detection():
    assert detect_language("Zwiebel schälen und in den Mixtopf geben") == "de"
    assert detect_language("Ajouter les oignons et faire revenir") == "fr"
    assert detect_language("Add the onions and cook") == "en"
    assert detect_language("قطع البصل واطبخه") == "ar"


def test_ojja_rules():
    r = convert_structured(OJJA)
    t = texts(r)
    assert r.language == "fr" and r.servings == 4
    assert t[0].startswith("Mettre ail dans le bol") and "5 sec/vitesse 5" in t[0]
    assert "faire revenir 3 min/120°C/vitesse 1" in t[1]
    assert "poivrons verts et tomates" in t[2] and "15 min/100°C" in t[2]
    assert "merguez" in t[3] and "sens inverse/vitesse mijotage" in t[3]
    assert "4 min/90°C" in t[4]
    assert r.steps[-1].off_machine
    assert "sans IA" in r.warnings[0]


def test_structure_and_convert_text():
    data = structure_text(SOUP_TEXT)
    assert data["title"] == "Linsensuppe" and data["yields"] == "4"
    assert len(data["ingredients"]) == 7 and len(data["instructions"]) == 4
    assert data["instructions"][0].endswith("andünsten.")  # wrapped line re-joined
    r = convert_structured(data)
    t = texts(r)
    assert "Petersilie in den trockenen Mixtopf" in t[0]           # dry bowl first
    assert "Zwiebel und Knoblauchzehen" in t[1]
    assert any("pürieren" in x and "Stufe 10" in x for x in t)
    assert any("bestreuen" in x for x in t)
    assert any("20 Min./100°C" in x for x in t)


def test_oven_and_deep_fry_stay_outside():
    r = convert_structured({"title": "Brik", "ingredients": ["4 feuilles de brick", "4 oeufs", "1 l d'huile"],
                            "instructions": ["Frire les bricks dans l'huile chaude.", "Enfourner 10 minutes à 200°C."]})
    assert all(s.off_machine for s in r.steps)


def test_long_cooking_split_and_capacity():
    r = convert_structured({"title": "Mloukhia", "ingredients": ["400 g Rindfleisch", "2,5 l Wasser"],
                            "instructions": ["Rindfleisch und Wasser zugeben und 3 Stunden köcheln lassen."]})
    assert sum(1 for s in r.steps if s.setting) == 2
    assert any("2,2 L" in w for w in r.warnings)


def test_docx_with_picture(tmp_path):
    import docx
    from PIL import Image
    d = docx.Document()
    d.add_heading("Lablabi", level=1)
    d.add_paragraph("Für 4 Personen")
    d.add_paragraph("Zutaten")
    for x in ["250 g Kichererbsen", "5 Knoblauchzehen", "1 TL Kreuzkümmel", "1,2 l Wasser"]:
        d.add_paragraph(x, style="List Bullet")
    d.add_paragraph("Zubereitung")
    d.add_paragraph("Knoblauch hacken.")
    d.add_paragraph("Kichererbsen und Wasser zugeben und 60 Minuten kochen.")
    img = io.BytesIO()
    Image.new("RGB", (400, 300), (200, 80, 40)).save(img, "JPEG")
    img.seek(0)
    d.add_picture(img)
    buf = io.BytesIO()
    d.save(buf)
    doc = read_document(buf.getvalue(), "lablabi.docx")
    assert doc.title == "Lablabi" and doc.images
    r = convert_structured(structure_text(doc.text, doc.title))
    assert r.title == "Lablabi" and len(r.ingredients) == 4
    assert any("60 Min./100°C/Linkslauf" in x for x in texts(r))


def test_pdf_text(tmp_path):
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    y = 800
    for line in SOUP_TEXT.splitlines():
        c.drawString(60, y, line.replace("½", "1/2"))
        y -= 16
    c.save()
    doc = read_document(buf.getvalue(), "suppe.pdf")
    data = structure_text(doc.text)
    assert data["title"] == "Linsensuppe" and len(data["instructions"]) >= 3


needs_tesseract = pytest.mark.skipif(shutil.which("tesseract") is None, reason="tesseract not installed")


def _recipe_card_png() -> bytes:
    from PIL import Image, ImageDraw, ImageFont
    lines = ["Banana Bread", "Serves 8", "Ingredients", "3 ripe bananas", "2 cups flour", "1/2 cup sugar",
             "1 egg", "Method", "1. Mash the bananas and mix with sugar and egg.", "2. Fold in the flour.",
             "3. Bake for 60 minutes at 175 C."]
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 34)
    except OSError:
        font = ImageFont.load_default()
    img = Image.new("RGB", (1300, 60 + 52 * len(lines)), "white")
    d = ImageDraw.Draw(img)
    for n, ln in enumerate(lines):
        d.text((40, 30 + n * 52), ln, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


@needs_tesseract
def test_ocr_photo_of_recipe():
    doc = read_document(_recipe_card_png(), "card.png")
    assert doc.ocr and "bananas" in doc.text.lower()
    data = structure_text(doc.text)
    r = convert_structured(data)
    assert r.title.lower().startswith("banana") and r.servings == 8
    assert any("mash" in x.lower() for x in texts(r))
    assert r.steps[-1].off_machine


@needs_tesseract
def test_scanned_pdf_uses_ocr():
    from PIL import Image
    png = Image.open(io.BytesIO(_recipe_card_png())).convert("RGB")
    buf = io.BytesIO()
    png.save(buf, "PDF", resolution=150)
    doc = read_document(buf.getvalue(), "scan.pdf")
    assert doc.ocr and "flour" in doc.text.lower()


def test_pipeline_url_with_picture(tmp_path, monkeypatch):
    """Recipe page with schema.org + og:image -> rules -> saved with local picture (no AI, no network)."""
    from PIL import Image
    from tunimix.library import Library
    from tunimix.pipeline import Options, process
    page = """<html><head><title>Ojja</title><script type="application/ld+json">{"@context":"https://schema.org","@type":"Recipe",
    "name":"Ojja aux merguez","recipeYield":"4","image":"https://img.example.tn/ojja.jpg",
    "recipeIngredient":["300 g de merguez","4 oeufs","3 gousses d'ail","400 g de tomates"],
    "recipeInstructions":[{"@type":"HowToStep","text":"Hacher l'ail."},{"@type":"HowToStep","text":"Ajouter les tomates et laisser mijoter 15 minutes."}]}
    </script></head><body></body></html>"""
    jpg = io.BytesIO()
    Image.new("RGB", (800, 600), (180, 60, 30)).save(jpg, "JPEG", quality=95)

    class Resp:
        def __init__(self, body, ctype):
            self.content, self.text, self.headers = body, body.decode("latin-1") if isinstance(body, bytes) else body, {"content-type": ctype}
            self.encoding = self.apparent_encoding = "utf-8"
            self.status_code = 200

        def raise_for_status(self):
            pass

    def get(url, headers=None, timeout=None):
        if url.endswith(".jpg"):
            return Resp(jpg.getvalue(), "image/jpeg")
        return Resp(page, "text/html")

    monkeypatch.setattr("requests.get", get)
    lib = Library(tmp_path)
    r = process("https://www.example.tn/ojja", Options(engine="rules"), lib)
    assert r.title == "Ojja aux merguez" and r.image_file.endswith("ojja-aux-merguez.jpg")
    assert (tmp_path / "images" / "ojja-aux-merguez.jpg").exists()
    assert lib.get("ojja-aux-merguez").source_type == "rules"
    from tunimix.outputs.pdf import build_pdf
    pdf = build_pdf([r], tmp_path / "x.pdf")
    assert pdf.stat().st_size > 20000  # picture embedded


def test_dish_name_without_ai_needs_page(monkeypatch, tmp_path):
    from tunimix.library import Library
    from tunimix.extract import search
    from tunimix.pipeline import Options, process
    monkeypatch.setattr(search, "ddg_search", lambda q, limit=6: [])
    with pytest.raises(RuntimeError, match="No recipe page"):
        process("kafteji", Options(engine="rules"), Library(tmp_path))
